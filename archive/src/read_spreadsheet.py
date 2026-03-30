import os

from openpyxl import load_workbook

path = os.path.expanduser("~/Documents/misc/software_inventory.xlsx")
workbook = load_workbook(filename=path)
sheet = workbook["SHH Applications"]

rows = sheet.iter_rows(values_only=True)

# Grab the header row and map column names to their index
headers = {name: i for i, name in enumerate(next(rows))}

for row in rows:
    if row[0] is None:
        break
    vendor = row[headers["Vendor Name"]]
    product = row[headers["Product Name"]]
    print(vendor, product)
