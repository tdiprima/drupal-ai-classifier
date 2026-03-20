#!/usr/bin/env python3
"""
audit_importer.py

Diagnostic audit script: determines why Drupal fields were not populated
during the Excel → Drupal import performed by drupal_importer.py.

For each unique product (vendor+product key), this script:
  1. Reads the same Excel spreadsheet as the importer.
  2. Captures raw per-sheet values BEFORE merging (to detect conflicts).
  3. Fetches the corresponding Drupal node via JSON:API.
  4. Compares each expected field value against what Drupal actually stored.
  5. Classifies any gap by root cause (missing source, mapping issue, etc.).

Output: a CSV diagnostic report + summary statistics logged to stdout.

Usage:
    python audit_importer.py --spreadsheet /path/to/file.xlsx
    python audit_importer.py --spreadsheet /path/to/file.xlsx --output report.csv
    python audit_importer.py --spreadsheet /path/to/file.xlsx --limit 10 --debug

Required .env variables (same as drupal_importer.py):
    DRUPAL_BASE_URL, DRUPAL_USERNAME, DRUPAL_PASSWORD, DRUPAL_CONTENT_TYPE
"""

import csv
import logging
import os
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import requests
import urllib3
from dotenv import load_dotenv

# Reuse constants and normalization functions from the importer
sys.path.insert(0, str(Path(__file__).parent))
from drupal_importer import (DATE_FIELDS, DEFAULT_SPREADSHEET,
                             FIELD_ALLOWED_VALUES, FIELD_MAP, SITE_COLUMNS,
                             SKIP_SHEETS, clean, normalize_date_value,
                             normalize_list_value)

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger(__name__)

# Reverse map for human-readable reporting
FIELD_MAP_REVERSE: dict[str, str] = {v: k for k, v in FIELD_MAP.items()}

# Every Drupal field we expect the importer to have populated from Excel
ALL_AUDITABLE_FIELDS: list[str] = list(FIELD_MAP.values()) + ["field_sites_used"]

# Polite delay between Drupal API calls (seconds)
DRUPAL_FETCH_DELAY = 0.2

DEFAULT_OUTPUT = Path("audit_report.csv")

# ---------------------------------------------------------------------------
# Status labels used in the output report
# ---------------------------------------------------------------------------

STATUS_OK = "Imported correctly"
STATUS_MISSING_SOURCE = "Missing source data"
STATUS_MAPPING_ISSUE = "Mapping issue"
STATUS_CONFLICT = "Conflict detected"
STATUS_SKIPPED = "Import skipped"
STATUS_TRANSFORM_ERROR = "Data transformation error"
STATUS_NODE_NOT_FOUND = "Node not found in Drupal"
STATUS_UNKNOWN = "Unknown"


# ---------------------------------------------------------------------------
# Data structures
# ---------------------------------------------------------------------------

@dataclass
class FieldAuditRow:
    """One row in the audit CSV — represents one field for one node."""
    node_id: str
    drupal_field: str
    excel_column: str
    excel_value: str
    drupal_value: str
    status: str
    explanation: str


@dataclass
class RawFieldData:
    """
    All raw cell values seen for one Drupal field across every sheet,
    recorded BEFORE merging. Used to detect empty-source and conflicts.
    """
    # List of (sheet_name, raw_cell_value) pairs
    values: list[tuple[str, str]] = field(default_factory=list)
    # True if the expected Excel column header was found in at least one sheet
    column_found_in_sheet: bool = False


# ---------------------------------------------------------------------------
# Argument parsing
# ---------------------------------------------------------------------------

def parse_args() -> dict:
    """Parse command-line arguments from sys.argv."""
    args = sys.argv[1:]

    if "--help" in args or "-h" in args:
        print(__doc__)
        sys.exit(0)

    spreadsheet = DEFAULT_SPREADSHEET
    output = DEFAULT_OUTPUT
    limit = None
    debug = "--debug" in args

    if "--spreadsheet" in args:
        idx = args.index("--spreadsheet")
        if idx + 1 < len(args):
            spreadsheet = Path(args[idx + 1])

    if "--output" in args:
        idx = args.index("--output")
        if idx + 1 < len(args):
            output = Path(args[idx + 1])

    if "--limit" in args:
        idx = args.index("--limit")
        if idx + 1 < len(args):
            try:
                limit = int(args[idx + 1])
            except ValueError:
                logger.error("--limit requires an integer argument")
                sys.exit(1)

    return {
        "spreadsheet": spreadsheet,
        "output": output,
        "limit": limit,
        "debug": debug,
    }


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

def load_config() -> dict:
    """Load and validate required environment variables from .env."""
    load_dotenv()
    required = ["DRUPAL_BASE_URL", "DRUPAL_USERNAME", "DRUPAL_PASSWORD", "DRUPAL_CONTENT_TYPE"]
    missing = [k for k in required if not os.environ.get(k)]
    if missing:
        logger.error("Missing required environment variables: %s", ", ".join(missing))
        sys.exit(1)
    return {
        "base_url": os.environ["DRUPAL_BASE_URL"].rstrip("/"),
        "username": os.environ["DRUPAL_USERNAME"],
        "password": os.environ["DRUPAL_PASSWORD"],
        "content_type": os.environ["DRUPAL_CONTENT_TYPE"],
    }


def apply_importer_transform(drupal_field: str, raw_value: str) -> tuple[str | int | None, str | None]:
    """
    Apply the same transformation drupal_importer.py would apply to a raw Excel value.
    Returns (transformed_value, failure_reason).
    failure_reason is None when the transform succeeds; a description string when it fails.
    """
    if not raw_value:
        return None, "empty value"

    if drupal_field in DATE_FIELDS:
        normalized = normalize_date_value(raw_value)
        if normalized is None:
            return None, f"date parse failed for {raw_value!r}"
        return normalized, None

    if drupal_field in FIELD_ALLOWED_VALUES:
        normalized = normalize_list_value(drupal_field, raw_value)
        if normalized is None:
            allowed = FIELD_ALLOWED_VALUES[drupal_field]
            return None, f"value {raw_value!r} not in allowed set {allowed}"
        return normalized, None

    return raw_value, None


# ---------------------------------------------------------------------------
# Excel reading with pre-merge conflict tracking
# ---------------------------------------------------------------------------

def load_excel_with_audit_data(
    filepath: Path,
) -> tuple[list[dict], dict[tuple, dict[str, RawFieldData]]]:
    """
    Read all data sheets and return two structures:

    1. merged_records — one dict per unique (vendor, product) key, same result
       as drupal_importer.py load_and_merge(). This is what the importer used.

    2. raw_field_data_by_key — maps each (vendor_lower, product_lower) key to
       a per-field dict of RawFieldData. Each RawFieldData captures every raw
       cell value seen across all sheets BEFORE the merge. Used to determine:
         - Was the Excel cell empty? (all raw values empty)
         - Was there a conflict? (multiple different non-empty values)
         - Was the column header even present? (column_found_in_sheet flag)
    """
    wb = openpyxl.load_workbook(filepath, data_only=True)
    merged: dict[tuple, dict] = {}
    raw_field_data_by_key: dict[tuple, dict[str, RawFieldData]] = {}

    for sheet_name in wb.sheetnames:
        if sheet_name.strip() in SKIP_SHEETS:
            continue

        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            logger.warning("Sheet '%s': empty, skipping", sheet_name)
            continue

        header = [str(c).strip() if c else "" for c in rows[0]]
        header_lower = [h.lower() for h in header]

        vendor_col = next(
            (i for i, h in enumerate(header_lower) if "vendor" in h and "name" in h), None
        )
        product_col = next(
            (i for i, h in enumerate(header_lower) if h == "product name"), None
        )
        if vendor_col is None or product_col is None:
            logger.warning("Sheet '%s': missing vendor/product columns, skipping", sheet_name)
            continue

        # Which site columns exist in this sheet?
        site_col_indices = {
            site: header.index(site)
            for site in SITE_COLUMNS
            if site in header
        }

        # Which FIELD_MAP columns exist in this sheet?
        field_col_indices: dict[str, int] = {}
        missing_from_header: set[str] = set()
        for col_label, drupal_field in FIELD_MAP.items():
            if col_label in header_lower:
                field_col_indices[drupal_field] = header_lower.index(col_label)
            else:
                missing_from_header.add(drupal_field)

        if missing_from_header:
            logger.debug(
                "Sheet '%s': Excel columns not found for Drupal fields: %s",
                sheet_name,
                ", ".join(sorted(missing_from_header)),
            )

        logger.info("Sheet '%s': %d data rows", sheet_name.strip(), len(rows) - 1)

        for row in rows[1:]:
            if len(row) <= max(vendor_col, product_col):
                continue

            vendor = clean(row[vendor_col])
            product = clean(row[product_col])

            desc_col = field_col_indices.get("field_description")
            description = (
                clean(row[desc_col])
                if (desc_col is not None and desc_col < len(row))
                else ""
            )

            # Rows without all three required fields are skipped by the importer
            if not vendor or not product or not description:
                missing = [
                    label
                    for label, val in [("vendor", vendor), ("product", product), ("description", description)]
                    if not val
                ]
                logger.debug(
                    "Sheet '%s': skipping row — missing required field(s): %s",
                    sheet_name,
                    ", ".join(missing),
                )
                continue

            key = (vendor.lower(), product.lower())

            # Collect site codes marked with "X" in this row
            sites = {
                site
                for site, col_idx in site_col_indices.items()
                if col_idx < len(row) and clean(row[col_idx]).upper() == "X"
            }

            # --- Initialize tracking structures for new keys ---
            if key not in raw_field_data_by_key:
                raw_field_data_by_key[key] = {}

            # Record raw cell value for every mapped field present in this sheet
            for drupal_field, col_idx in field_col_indices.items():
                raw_val = clean(row[col_idx]) if col_idx < len(row) else ""
                if drupal_field not in raw_field_data_by_key[key]:
                    raw_field_data_by_key[key][drupal_field] = RawFieldData(column_found_in_sheet=True)
                rfd = raw_field_data_by_key[key][drupal_field]
                rfd.column_found_in_sheet = True
                rfd.values.append((sheet_name, raw_val))

            # Record that certain fields had no corresponding column in this sheet
            for drupal_field in missing_from_header:
                if drupal_field not in raw_field_data_by_key[key]:
                    raw_field_data_by_key[key][drupal_field] = RawFieldData(column_found_in_sheet=False)

            # Record raw sites data
            if "field_sites_used" not in raw_field_data_by_key[key]:
                raw_field_data_by_key[key]["field_sites_used"] = RawFieldData(
                    column_found_in_sheet=bool(site_col_indices)
                )
            sites_rfd = raw_field_data_by_key[key]["field_sites_used"]
            sites_rfd.column_found_in_sheet = sites_rfd.column_found_in_sheet or bool(site_col_indices)
            sites_rfd.values.append((sheet_name, ",".join(sorted(sites))))

            # --- Merge record (mirrors drupal_importer.py exactly) ---
            if key in merged:
                existing = merged[key]
                existing["field_sites_used"] |= sites
                for drupal_field, col_idx in field_col_indices.items():
                    if drupal_field in ("field_vendor_name", "field_product_name"):
                        continue
                    new_val = clean(row[col_idx]) if col_idx < len(row) else ""
                    old_val = existing.get(drupal_field, "")
                    if new_val and not old_val:
                        # Fill in previously-empty field with new value
                        existing[drupal_field] = new_val
                    elif new_val and old_val and new_val != old_val:
                        # Conflict: importer keeps the first value; we log for the audit
                        logger.debug(
                            "Conflict '%s — %s' [%s]: keeping '%s', ignoring '%s'",
                            vendor, product, drupal_field, old_val, new_val,
                        )
            else:
                record: dict = {
                    "field_vendor_name": vendor,
                    "field_product_name": product,
                    "field_sites_used": sites,
                }
                for drupal_field, col_idx in field_col_indices.items():
                    val = clean(row[col_idx]) if col_idx < len(row) else ""
                    if val:
                        record[drupal_field] = val
                merged[key] = record

    logger.info("Total unique products after merging: %d", len(merged))
    return list(merged.values()), raw_field_data_by_key


# ---------------------------------------------------------------------------
# Drupal node fetching
# ---------------------------------------------------------------------------

def build_session(config: dict) -> requests.Session:
    """Create an authenticated HTTP session for JSON:API requests."""
    session = requests.Session()
    session.auth = (config["username"], config["password"])
    session.headers.update({
        "Accept": "application/vnd.api+json",
        "Content-Type": "application/vnd.api+json",
    })
    # SSL verification intentionally disabled for internal university server
    session.verify = False
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
    return session


def fetch_drupal_node(
    session: requests.Session,
    base_url: str,
    content_type: str,
    title: str,
) -> dict | None:
    """
    Fetch a Drupal node by its title via JSON:API filter query.
    Returns the node's attributes dict (including a synthetic '_uuid' key),
    or None if the node cannot be found or the request fails.
    """
    url = f"{base_url}/jsonapi/node/{content_type}"
    fields_param = ",".join(ALL_AUDITABLE_FIELDS + ["title", "drupal_internal__nid"])
    params = {
        "filter[title]": title,
        f"fields[node--{content_type}]": fields_param,
    }
    try:
        response = session.get(url, params=params, timeout=30)
        response.raise_for_status()
        data = response.json()
        nodes = data.get("data", [])

        if not nodes:
            logger.warning("No Drupal node found for title: %r", title)
            return None

        if len(nodes) > 1:
            logger.warning("Multiple nodes found for title %r — using first result", title)

        node = nodes[0]
        attrs = node.get("attributes", {})
        attrs["_uuid"] = node.get("id", "")
        return attrs

    except requests.exceptions.RequestException as exc:
        logger.error("Failed to fetch Drupal node %r: %s", title, exc)
        return None


def extract_drupal_field_value(attrs: dict, drupal_field: str) -> str:
    """
    Extract a normalized string value for one field from a Drupal node attributes dict.
    Multi-value list fields (e.g. field_sites_used) are returned as comma-separated strings.
    """
    value = attrs.get(drupal_field)
    if value is None:
        return ""
    if isinstance(value, list):
        # JSON:API returns multi-value fields as [{"value": "..."}, ...]
        parts = [
            str(item.get("value", item)) if isinstance(item, dict) else str(item)
            for item in value
        ]
        return ",".join(sorted(parts))
    if isinstance(value, dict):
        return str(value.get("value", "")).strip()
    return str(value).strip()


# ---------------------------------------------------------------------------
# Gap classification
# ---------------------------------------------------------------------------

def classify_regular_field(
    drupal_field: str,
    excel_merged_value: str,
    drupal_value: str,
    rfd: RawFieldData | None,
) -> tuple[str, str]:
    """
    Determine why a regular (non-sites) field is empty in Drupal.
    Works through the possible causes in priority order.
    Returns (status_label, explanation_string).
    """
    # --- Drupal has a value: check for match or acceptable divergence ---
    if drupal_value:
        if not excel_merged_value:
            return (
                STATUS_OK,
                f"Drupal has value {drupal_value!r}; Excel source was empty "
                "(node may have been patched or populated via another path)",
            )

        transformed, _ = apply_importer_transform(drupal_field, excel_merged_value)
        drupal_lower = drupal_value.lower().strip()
        excel_lower = str(transformed).lower().strip() if transformed is not None else ""

        if drupal_lower == excel_lower or drupal_value == excel_merged_value:
            return STATUS_OK, "Values match"

        return (
            STATUS_OK,
            f"Drupal has {drupal_value!r}; Excel source was {excel_merged_value!r} "
            "— may differ due to fuzzy match, post-import patch, or conflict resolution",
        )

    # --- Drupal field is empty from this point ---

    # Cause 1: The expected Excel column was never found in any sheet header
    if rfd is None or not rfd.column_found_in_sheet:
        excel_col = FIELD_MAP_REVERSE.get(drupal_field, drupal_field)
        return (
            STATUS_MAPPING_ISSUE,
            f"Excel column {excel_col!r} not found in any sheet header — "
            "field was never read from the spreadsheet",
        )

    # Cause 2: Every raw cell value across all sheets was empty
    non_empty_raws = [(sheet, val) for sheet, val in rfd.values if val]
    if not non_empty_raws:
        return (
            STATUS_MISSING_SOURCE,
            "Excel cell was empty in every sheet — no source data existed to import",
        )

    # Cause 3: Multiple sheets had different non-empty values (conflict);
    # importer keeps the first, but the first may itself have had issues
    unique_values = {val for _, val in non_empty_raws}
    if len(unique_values) > 1:
        conflict_summary = "; ".join(
            f"{sheet}: {val!r}" for sheet, val in non_empty_raws
        )
        return (
            STATUS_CONFLICT,
            f"Conflicting values across sheets: {conflict_summary} — "
            "importer kept the first non-empty value; Drupal field is still empty, "
            "suggesting the winning value also failed validation",
        )

    # Cause 4: The single raw value failed the importer's transformation
    raw_value = list(unique_values)[0]
    transformed, failure_reason = apply_importer_transform(drupal_field, raw_value)
    if transformed is None:
        return (
            STATUS_TRANSFORM_ERROR,
            f"Value {raw_value!r} failed transformation: {failure_reason} — "
            "importer would have logged a warning and skipped this field",
        )

    # Cause 5: Value passed validation but Drupal field is still empty
    # Most likely a Drupal-side rejection (HTTP 422) or silent importer bug
    return (
        STATUS_SKIPPED,
        f"Excel had valid value {raw_value!r} but Drupal field is empty — "
        "possible causes: Drupal rejected it via HTTP 422, the field machine name "
        "differed from what the importer sent, or a concurrent import overwrote it",
    )


def classify_sites_field(
    excel_merged_sites: set,
    drupal_value: str,
    rfd: RawFieldData | None,
) -> tuple[str, str]:
    """Classify the field_sites_used multi-value field."""
    excel_sites_str = ",".join(sorted(s.lower() for s in excel_merged_sites))

    if drupal_value and excel_sites_str:
        drupal_set = set(drupal_value.lower().split(","))
        excel_set = set(excel_sites_str.split(",")) - {""}
        if drupal_set == excel_set:
            return STATUS_OK, "Site values match"
        return (
            STATUS_OK,
            f"Drupal has {drupal_value!r}; Excel has {excel_sites_str!r} — "
            "difference may reflect union across sheets or post-import patch",
        )

    if drupal_value and not excel_sites_str:
        return (
            STATUS_OK,
            "Drupal has site values; Excel had no X-marked site columns for this product",
        )

    if not excel_sites_str:
        return (
            STATUS_MISSING_SOURCE,
            "No site columns (SBUH, SBSH, SBELIH, etc.) were marked with X in Excel — "
            "no site data existed to import",
        )

    if rfd and not rfd.column_found_in_sheet:
        return (
            STATUS_MAPPING_ISSUE,
            "Site columns not found in any sheet header — field_sites_used could not be populated",
        )

    return (
        STATUS_SKIPPED,
        f"Excel had sites {excel_sites_str!r} but Drupal field is empty — "
        "possible Drupal-side rejection or field machine name mismatch",
    )


# ---------------------------------------------------------------------------
# Audit one record
# ---------------------------------------------------------------------------

def audit_record(
    record: dict,
    rfd_for_record: dict[str, RawFieldData],
    drupal_attrs: dict | None,
    node_title: str,
) -> list[FieldAuditRow]:
    """
    Compare one merged Excel record against its Drupal node attributes.
    Returns one FieldAuditRow per auditable field.
    """
    audit_rows: list[FieldAuditRow] = []

    for drupal_field in ALL_AUDITABLE_FIELDS:
        excel_col = FIELD_MAP_REVERSE.get(
            drupal_field,
            "site columns (SBUH/SBSH/…)" if drupal_field == "field_sites_used" else "?",
        )
        rfd = rfd_for_record.get(drupal_field)

        # Cannot compare if the Drupal node was not found
        if drupal_attrs is None:
            audit_rows.append(FieldAuditRow(
                node_id=node_title,
                drupal_field=drupal_field,
                excel_column=excel_col,
                excel_value=str(record.get(drupal_field, "")),
                drupal_value="",
                status=STATUS_NODE_NOT_FOUND,
                explanation="Drupal node could not be retrieved — no comparison possible",
            ))
            continue

        drupal_value = extract_drupal_field_value(drupal_attrs, drupal_field)

        if drupal_field == "field_sites_used":
            excel_sites = record.get("field_sites_used", set())
            excel_value_str = ",".join(sorted(s.lower() for s in excel_sites))
            status, explanation = classify_sites_field(excel_sites, drupal_value, rfd)
        else:
            excel_value_str = str(record.get(drupal_field, ""))
            status, explanation = classify_regular_field(
                drupal_field, excel_value_str, drupal_value, rfd
            )

        audit_rows.append(FieldAuditRow(
            node_id=node_title,
            drupal_field=drupal_field,
            excel_column=excel_col,
            excel_value=excel_value_str,
            drupal_value=drupal_value,
            status=status,
            explanation=explanation,
        ))

    return audit_rows


# ---------------------------------------------------------------------------
# Report output
# ---------------------------------------------------------------------------

def write_csv_report(audit_rows: list[FieldAuditRow], output_path: Path) -> None:
    """Write all audit rows to a CSV file."""
    fieldnames = [
        "node_id",
        "drupal_field",
        "excel_column",
        "excel_value",
        "drupal_value",
        "status",
        "explanation",
    ]
    with output_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        for row in audit_rows:
            writer.writerow({
                "node_id": row.node_id,
                "drupal_field": row.drupal_field,
                "excel_column": row.excel_column,
                "excel_value": row.excel_value,
                "drupal_value": row.drupal_value,
                "status": row.status,
                "explanation": row.explanation,
            })
    logger.info("Audit report written: %s (%d rows)", output_path, len(audit_rows))


def print_summary(audit_rows: list[FieldAuditRow], total_nodes: int) -> None:
    """Log summary statistics and the top offending fields."""
    total_fields = len(audit_rows)
    by_status: dict[str, int] = {}
    for row in audit_rows:
        by_status[row.status] = by_status.get(row.status, 0) + 1

    sep = "=" * 64
    logger.info(sep)
    logger.info("AUDIT SUMMARY")
    logger.info("  Total nodes analyzed           : %d", total_nodes)
    logger.info("  Total field checks             : %d", total_fields)
    logger.info("  %-38s : %d", STATUS_OK, by_status.get(STATUS_OK, 0))
    logger.info("  %-38s : %d", STATUS_MISSING_SOURCE, by_status.get(STATUS_MISSING_SOURCE, 0))
    logger.info("  %-38s : %d", STATUS_MAPPING_ISSUE, by_status.get(STATUS_MAPPING_ISSUE, 0))
    logger.info("  %-38s : %d", STATUS_CONFLICT, by_status.get(STATUS_CONFLICT, 0))
    logger.info("  %-38s : %d", STATUS_SKIPPED, by_status.get(STATUS_SKIPPED, 0))
    logger.info("  %-38s : %d", STATUS_TRANSFORM_ERROR, by_status.get(STATUS_TRANSFORM_ERROR, 0))
    logger.info("  %-38s : %d", STATUS_NODE_NOT_FOUND, by_status.get(STATUS_NODE_NOT_FOUND, 0))
    logger.info("  %-38s : %d", STATUS_UNKNOWN, by_status.get(STATUS_UNKNOWN, 0))
    logger.info(sep)

    # Fields with the most issues (excluding node-not-found, which is a fetch error)
    field_issue_counts: dict[str, int] = {}
    for row in audit_rows:
        if row.status not in (STATUS_OK, STATUS_NODE_NOT_FOUND):
            field_issue_counts[row.drupal_field] = field_issue_counts.get(row.drupal_field, 0) + 1

    if field_issue_counts:
        logger.info("Top fields with issues (nodes affected):")
        ranked = sorted(field_issue_counts.items(), key=lambda x: -x[1])
        for fname, count in ranked[:10]:
            logger.info("  %-48s : %d", fname, count)
        logger.info(sep)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    args = parse_args()

    if args["debug"]:
        logging.getLogger().setLevel(logging.DEBUG)

    spreadsheet: Path = args["spreadsheet"]
    output: Path = args["output"]
    limit: int | None = args["limit"]

    if not spreadsheet.exists():
        logger.error("Spreadsheet not found: %s", spreadsheet)
        sys.exit(1)

    config = load_config()

    logger.info("Loading spreadsheet: %s", spreadsheet)
    records, raw_field_data_by_key = load_excel_with_audit_data(spreadsheet)

    if limit:
        records = records[:limit]
        logger.info("Limiting to first %d records", limit)

    total_nodes = len(records)
    logger.info("Connecting to Drupal: %s", config["base_url"])
    session = build_session(config)

    all_audit_rows: list[FieldAuditRow] = []

    for idx, record in enumerate(records, 1):
        vendor = record.get("field_vendor_name", "")
        product = record.get("field_product_name", "")
        node_title = f"{vendor} — {product}"
        key = (vendor.lower(), product.lower())

        logger.info("[%d/%d] Auditing: %s", idx, total_nodes, node_title)

        drupal_attrs = fetch_drupal_node(
            session, config["base_url"], config["content_type"], node_title
        )
        rfd_for_record = raw_field_data_by_key.get(key, {})

        audit_rows = audit_record(record, rfd_for_record, drupal_attrs, node_title)
        all_audit_rows.extend(audit_rows)

        # Avoid hammering the Drupal server
        if idx < total_nodes:
            time.sleep(DRUPAL_FETCH_DELAY)

    write_csv_report(all_audit_rows, output)
    print_summary(all_audit_rows, total_nodes)


if __name__ == "__main__":
    main()
