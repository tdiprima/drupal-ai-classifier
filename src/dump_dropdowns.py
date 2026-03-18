#!/usr/bin/env python3
"""
dump_dropdowns.py

Prints the contents of the Dropdowns sheet from the spreadsheet
so you can see the allowed values for each list field.

Usage:
    python dump_dropdowns.py
    python dump_dropdowns.py /path/to/file.xlsx
"""

import sys
from pathlib import Path

import openpyxl

DEFAULT_SPREADSHEET = Path.home() / "Documents" / "misc" / "software_inventory.xlsx"


def main() -> None:
    filepath = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SPREADSHEET

    if not filepath.exists():
        print(f"ERROR: file not found: {filepath}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(filepath, data_only=True)

    if "Dropdowns" not in wb.sheetnames:
        print("No 'Dropdowns' sheet found. Available sheets:")
        for name in wb.sheetnames:
            print(f"  {name}")
        sys.exit(0)

    ws = wb["Dropdowns"]
    for row in ws.iter_rows(values_only=True):
        if any(cell is not None for cell in row):
            print("\t".join("" if cell is None else str(cell) for cell in row))


if __name__ == "__main__":
    main()
