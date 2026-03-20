# The whole point is scanning the inventory workbook for Microsoft SB Applications.
# Dumps the whole Excel sheet into memory.  See scan_inventory_streaming.py
import openpyxl
from pathlib import Path


def find_sb_applications():
    workbook_path = Path.home() / "Documents/misc/software_inventory.xlsx"

    wb = openpyxl.load_workbook(
        workbook_path,
        data_only=True,
        read_only=True
    )

    skip_sheets = {
        "MASTER Spreadsheet",
        "Priority Level Definitions",
        "Dropdowns",
        "Sheet1"
    }

    for sheet_name in wb.sheetnames:
        if sheet_name.strip() in skip_sheets:
            continue

        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            continue

        header = [str(c).strip().lower() if c else "" for c in rows[0]]

        vendor_col = next(
            (i for i, h in enumerate(header) if "vendor" in h and "name" in h),
            None
        )

        product_col = next(
            (i for i, h in enumerate(header) if h == "product name"),
            None
        )

        criticality_col = next(
            (i for i, h in enumerate(header) if h == "business criticality level"),
            None
        )

        if vendor_col is None or product_col is None:
            continue

        for row_num, row in enumerate(rows[1:], start=2):
            if len(row) <= max(vendor_col, product_col):
                continue

            vendor = str(row[vendor_col] or "").strip()
            product = str(row[product_col] or "").strip()
            crit = (
                str(row[criticality_col] or "").strip()
                if criticality_col is not None
                else ""
            )

            if vendor.lower() == "microsoft" and "sb application" in product.lower():
                print(
                    f"Sheet: {sheet_name!r}, Row {row_num}: "
                    f"vendor={vendor!r}, product={product!r}, criticality={crit!r}"
                )

    wb.close()


if __name__ == "__main__":
    find_sb_applications()
