import os

from openpyxl import load_workbook

path = os.path.expanduser("~/Documents/misc/software_inventory.xlsx")
workbook = load_workbook(filename=path)
sheet = workbook["SHH Applications"]
# sheet = workbook.worksheets[1]  # Gets the second sheet

print(sheet["A1"].value)  # Output: 'Division'

for row in sheet.iter_rows(values_only=True):
    if all(cell is None for cell in row):
        break
    print(row)
