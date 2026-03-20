#!/usr/bin/env python3
"""
import_no_description.py

Imports Drupal nodes for spreadsheet rows that have Vendor Name and Product Name
but a blank Description field — rows that drupal_importer.py currently skips.

Supports resuming interrupted imports via a progress file.

Usage:
    python import_no_description.py --dry-run               # preview without posting
    python import_no_description.py                         # live import (or resume)
    python import_no_description.py --limit 10              # import first N records only
    python import_no_description.py --reset                 # clear progress and start over
    python import_no_description.py --spreadsheet /path/to/file.xlsx
    python import_no_description.py --help

Required .env variables:
    DRUPAL_BASE_URL       e.g. http://bmi-capella.uhmc.sunysb.edu
    DRUPAL_USERNAME       dedicated scanner account
    DRUPAL_PASSWORD       scanner account password
    DRUPAL_CONTENT_TYPE   machine name of the content type
"""

import logging
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl
import requests
import urllib3

from drupal_importer import (
    DEFAULT_SPREADSHEET,
    DATE_FIELDS,
    FIELD_ALLOWED_VALUES,
    FIELD_INTEGER_FIELDS,
    FIELD_MAP,
    MAX_RETRIES,
    RETRY_DELAY_SECONDS,
    SITE_COLUMNS,
    SKIP_SHEETS,
    build_payload,
    clean,
    extract_invalid_fields,
    load_config,
    load_progress,
    mark_complete,
    normalize_date_value,
    normalize_list_value,
    post_node_with_retry,
    save_progress,
)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger(__name__)

DEFAULT_PROGRESS_FILE = Path(__file__).parent.parent / "import_nodesc_progress.json"


# ---------------------------------------------------------------------------
# Argument parsing
# ---------------------------------------------------------------------------

def parse_args() -> dict:
    """Parse command-line arguments from sys.argv."""
    args = sys.argv[1:]

    if "--help" in args or "-h" in args:
        print(__doc__)
        sys.exit(0)

    dry_run = "--dry-run" in args
    debug = "--debug" in args
    reset = "--reset" in args
    limit = None
    spreadsheet = DEFAULT_SPREADSHEET
    progress_file = DEFAULT_PROGRESS_FILE

    if "--limit" in args:
        idx = args.index("--limit")
        if idx + 1 < len(args):
            try:
                limit = int(args[idx + 1])
            except ValueError:
                logger.error("--limit requires an integer argument")
                sys.exit(1)

    if "--spreadsheet" in args:
        idx = args.index("--spreadsheet")
        if idx + 1 < len(args):
            spreadsheet = Path(args[idx + 1])

    if "--progress-file" in args:
        idx = args.index("--progress-file")
        if idx + 1 < len(args):
            progress_file = Path(args[idx + 1])

    return {
        "dry_run": dry_run,
        "debug": debug,
        "reset": reset,
        "limit": limit,
        "spreadsheet": spreadsheet,
        "progress_file": progress_file,
    }


# ---------------------------------------------------------------------------
# Spreadsheet loading
# ---------------------------------------------------------------------------

def load_no_description_records(filepath: Path) -> list[dict]:
    """
    Read all data sheets and return merged records for rows that have
    Vendor Name and Product Name but a blank Description.

    Rows are merged by vendor+product key the same way as drupal_importer.py.
    """
    wb = openpyxl.load_workbook(filepath, data_only=True)
    merged: dict[tuple, dict] = {}

    for sheet_name in wb.sheetnames:
        if sheet_name.strip() in SKIP_SHEETS:
            continue

        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            logger.warning("Sheet '%s' is empty, skipping", sheet_name)
            continue

        header = [str(c).strip() if c else "" for c in rows[0]]
        header_lower = [h.lower() for h in header]

        vendor_col = next(
            (i for i, h in enumerate(header_lower) if "vendor" in h and "name" in h), None
        )
        product_col = next(
            (i for i, h in enumerate(header_lower) if h == "product name"), None
        )
        if vendor_col is None or product_col is None:
            logger.warning("Sheet '%s' missing vendor/product columns, skipping", sheet_name)
            continue

        site_col_indices = {
            site: header.index(site)
            for site in SITE_COLUMNS
            if site in header
        }

        field_col_indices = {
            drupal_field: header_lower.index(col_label)
            for col_label, drupal_field in FIELD_MAP.items()
            if col_label in header_lower
        }

        desc_col = field_col_indices.get("field_description")
        found_in_sheet = 0

        for row in rows[1:]:
            if len(row) <= max(vendor_col, product_col):
                continue

            vendor = clean(row[vendor_col])
            product = clean(row[product_col])

            if not vendor or not product:
                continue

            description = clean(row[desc_col]) if (desc_col is not None and desc_col < len(row)) else ""
            if description:
                # Has a description — handled by the main importer
                continue

            key = (vendor.lower(), product.lower())

            sites = {
                site
                for site, col_idx in site_col_indices.items()
                if col_idx < len(row) and clean(row[col_idx]).upper() == "X"
            }

            if key in merged:
                existing = merged[key]
                existing["field_sites_used"] |= sites
                for drupal_field, col_idx in field_col_indices.items():
                    if drupal_field in ("field_vendor_name", "field_product_name", "field_description"):
                        continue
                    new_val = clean(row[col_idx]) if col_idx < len(row) else ""
                    old_val = existing.get(drupal_field, "")
                    if new_val and old_val and new_val != old_val:
                        logger.warning(
                            "Conflict '%s — %s' [%s]: keeping '%s', ignoring '%s'",
                            vendor, product, drupal_field, old_val, new_val,
                        )
            else:
                record: dict = {
                    "field_vendor_name": vendor,
                    "field_product_name": product,
                    "field_sites_used": sites,
                }
                for drupal_field, col_idx in field_col_indices.items():
                    if drupal_field == "field_description":
                        continue
                    val = clean(row[col_idx]) if col_idx < len(row) else ""
                    if val:
                        record[drupal_field] = val
                merged[key] = record
                found_in_sheet += 1

        if found_in_sheet:
            logger.info("Sheet '%s': %d no-description row(s)", sheet_name.strip(), found_in_sheet)

    logger.info("Total unique no-description products: %d", len(merged))
    return list(merged.values())


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    args = parse_args()

    if args["debug"]:
        logging.getLogger().setLevel(logging.DEBUG)

    config = load_config()
    spreadsheet = args["spreadsheet"]
    progress_file = args["progress_file"]

    if not spreadsheet.exists():
        logger.error("Spreadsheet not found: %s", spreadsheet)
        sys.exit(1)

    if args["reset"]:
        if progress_file.exists():
            progress_file.unlink()
            logger.info("Progress file cleared — starting from scratch")
        else:
            logger.info("No progress file found — nothing to reset")

    logger.info("Loading spreadsheet: %s", spreadsheet)
    records = load_no_description_records(spreadsheet)

    if args["limit"]:
        records = records[: args["limit"]]
        logger.info("Limiting to first %d records", args["limit"])

    total = len(records)

    if args["dry_run"]:
        logger.info("DRY RUN — no nodes will be posted to Drupal")
        for record in records:
            sites = sorted(record.get("field_sites_used", set()))
            logger.info(
                "  %s — %s | sites: %s",
                record.get("field_vendor_name"),
                record.get("field_product_name"),
                sites,
            )
        logger.info("Would create %d node(s)", total)
        return

    completed = load_progress(progress_file)
    already_done = len(completed)
    if already_done:
        logger.info(
            "Resuming: %d record(s) already imported, %d remaining",
            already_done,
            total - already_done,
        )

    session = requests.Session()
    session.auth = (config["username"], config["password"])
    session.headers.update({
        "Accept": "application/vnd.api+json",
        "Content-Type": "application/vnd.api+json",
    })
    session.verify = False
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    succeeded = 0
    failed = 0
    skipped = 0
    start_time = time.monotonic()

    run_stats = {
        "started": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "succeeded": 0,
        "failed": 0,
        "skipped": 0,
    }

    try:
        for i, record in enumerate(records, 1):
            vendor = record.get("field_vendor_name", "")
            product = record.get("field_product_name", "")
            title = f"{vendor} \u2014 {product}"
            key = (vendor.lower(), product.lower())

            if key in completed:
                logger.debug("[%d/%d] SKIP (already imported): %s", i, total, title)
                skipped += 1
                continue

            logger.info("[%d/%d] %s", i, total, title)
            payload = build_payload(record, config["content_type"])

            if post_node_with_retry(session, config["base_url"], config["content_type"], payload):
                succeeded += 1
                run_stats.update({"succeeded": succeeded, "failed": failed, "skipped": skipped})
                mark_complete(progress_file, completed, key, run_stats)
            else:
                failed += 1
                run_stats.update({"succeeded": succeeded, "failed": failed, "skipped": skipped})
                logger.error("  FAILED — will retry on next run")

    except KeyboardInterrupt:
        logger.warning("Import interrupted by user (Ctrl+C)")

    finally:
        elapsed = time.monotonic() - start_time
        run_stats.update({
            "succeeded": succeeded,
            "failed": failed,
            "skipped": skipped,
            "finished": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "elapsed_seconds": round(elapsed, 1),
        })
        save_progress(progress_file, completed, run_stats)

        logger.info("=" * 60)
        logger.info("IMPORT (NO DESCRIPTION) SUMMARY")
        logger.info("  Total found          : %d", total)
        logger.info("  Skipped (done prior) : %d", skipped)
        logger.info("  Succeeded this run   : %d", succeeded)
        logger.info("  Failed this run      : %d", failed)
        logger.info("  Still to import      : %d", total - skipped - succeeded)
        logger.info("  Elapsed              : %.1fs", elapsed)
        logger.info("  Progress file        : %s", progress_file)
        if failed:
            logger.warning("Re-run this script to retry the %d failed record(s)", failed)


if __name__ == "__main__":
    main()
