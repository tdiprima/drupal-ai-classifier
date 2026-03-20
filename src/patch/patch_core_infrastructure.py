#!/usr/bin/env python3
"""
patch_core_infrastructure.py

Reads software_inventory.xlsx, finds all rows where 'Business Criticality Level'
is "Core Infrastructure", then patches field_business_criticality_level on the
matching Drupal nodes to "core_infrastructure".

Source rows are deduplicated by title before patching — a product appearing in
multiple sheets is only patched once.

Usage:
    python patch_core_infrastructure.py --dry-run       # preview without patching
    python patch_core_infrastructure.py                 # live patch
    python patch_core_infrastructure.py --help

Required .env variables: same as drupal_importer.py

Note: "core_infrastructure" is the Drupal machine key for "Core Infrastructure".
If Drupal rejects the patch with 422, verify the exact key with:
    python src/utility/fetch_allowed_values.py
"""

import logging
import os
import sys
import time
from pathlib import Path

import openpyxl
import requests
import urllib3
from dotenv import load_dotenv

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger(__name__)

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

DEFAULT_SPREADSHEET = Path.home() / "Documents" / "misc" / "software_inventory.xlsx"
PAGE_SIZE = 50
PATCH_VALUE = "core_infrastructure"

# Sheets that contain metadata, not software rows
SKIP_SHEETS = {"MASTER Spreadsheet", "Priority Level Definitions", "Dropdowns", "Sheet1"}

# The em dash used in Drupal node titles: "{Vendor} — {Product}"
EM_DASH = "\u2014"

# Column header (case-insensitive) we filter on
CRITICALITY_COLUMN = "business criticality level"
TARGET_VALUE = "core infrastructure"


def parse_args() -> dict:
    """Parse command-line arguments from sys.argv."""
    args = sys.argv[1:]

    if "--help" in args or "-h" in args:
        print(__doc__)
        sys.exit(0)

    spreadsheet = DEFAULT_SPREADSHEET
    if "--spreadsheet" in args:
        idx = args.index("--spreadsheet")
        if idx + 1 < len(args):
            spreadsheet = Path(args[idx + 1])

    return {
        "dry_run": "--dry-run" in args,
        "spreadsheet": spreadsheet,
    }


def iter_core_infrastructure_titles(spreadsheet_path: Path):
    """
    Yield unique Drupal node titles from the spreadsheet where
    'Business Criticality Level' == 'Core Infrastructure'.

    Iterates row-by-row (memory efficient — no dataframe).
    Deduplicates by title across sheets.
    """
    wb = openpyxl.load_workbook(spreadsheet_path, data_only=True, read_only=True)
    seen: set[str] = set()

    for sheet_name in wb.sheetnames:
        if sheet_name.strip() in SKIP_SHEETS:
            continue

        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)

        try:
            raw_header = next(rows)
        except StopIteration:
            logger.warning("Sheet '%s' is empty, skipping", sheet_name)
            continue

        header = [str(c).strip().lower() if c else "" for c in raw_header]

        vendor_col = next(
            (i for i, h in enumerate(header) if "vendor" in h and "name" in h), None
        )
        product_col = next(
            (i for i, h in enumerate(header) if h == "product name"), None
        )
        criticality_col = next(
            (i for i, h in enumerate(header) if h == CRITICALITY_COLUMN), None
        )

        if vendor_col is None or product_col is None or criticality_col is None:
            logger.warning(
                "Sheet '%s' missing required columns (vendor/product/criticality), skipping",
                sheet_name,
            )
            continue

        for row in rows:
            max_col = max(vendor_col, product_col, criticality_col)
            if len(row) <= max_col:
                continue

            criticality = str(row[criticality_col] or "").strip()
            if criticality.lower() != TARGET_VALUE:
                continue

            vendor = str(row[vendor_col] or "").strip()
            product = str(row[product_col] or "").strip()
            if not vendor or not product:
                continue

            title = f"{vendor} {EM_DASH} {product}"
            if title not in seen:
                seen.add(title)
                yield title

    wb.close()


def fetch_all_nodes(
    session: requests.Session, base_url: str, content_type: str
) -> dict[str, str]:
    """Fetch all nodes and return a {title: uuid} mapping."""
    index: dict[str, str] = {}
    url = f"{base_url}/jsonapi/node/{content_type}"
    params = {
        f"fields[node--{content_type}]": "title",
        "page[limit]": PAGE_SIZE,
    }

    page = 0
    while url:
        page += 1
        response = session.get(
            url, params=params if page == 1 else None, timeout=30
        )
        if response.status_code != 200:
            logger.error("Failed to fetch node list (HTTP %s)", response.status_code)
            sys.exit(1)

        data = response.json()
        for node in data.get("data", []):
            title = node.get("attributes", {}).get("title", "")
            uuid = node.get("id", "")
            if title and uuid:
                index[title] = uuid

        logger.info("  Fetched page %d (%d nodes so far)", page, len(index))
        url = data.get("links", {}).get("next", {}).get("href")
        params = None

    return index


def patch_node(
    session: requests.Session,
    base_url: str,
    content_type: str,
    uuid: str,
    title: str,
) -> bool:
    """PATCH field_business_criticality_level to PATCH_VALUE. Returns True on success."""
    url = f"{base_url}/jsonapi/node/{content_type}/{uuid}"
    payload = {
        "data": {
            "type": f"node--{content_type}",
            "id": uuid,
            "attributes": {
                "field_business_criticality_level": PATCH_VALUE,
            },
        }
    }
    try:
        response = session.patch(url, json=payload, timeout=30)
        response.raise_for_status()
        return True
    except requests.exceptions.HTTPError as exc:
        logger.error(
            "HTTP %s patching '%s': %s",
            exc.response.status_code,
            title,
            exc.response.text[:400],
        )
    except requests.exceptions.RequestException as exc:
        logger.error("Request error patching '%s': %s", title, exc)
    return False


def main() -> None:
    args = parse_args()

    load_dotenv()
    required = [
        "DRUPAL_BASE_URL",
        "DRUPAL_USERNAME",
        "DRUPAL_PASSWORD",
        "DRUPAL_CONTENT_TYPE",
    ]
    missing = [k for k in required if not os.environ.get(k)]
    if missing:
        logger.error("Missing required environment variables: %s", ", ".join(missing))
        sys.exit(1)

    base_url = os.environ["DRUPAL_BASE_URL"].rstrip("/")
    username = os.environ["DRUPAL_USERNAME"]
    password = os.environ["DRUPAL_PASSWORD"]
    content_type = os.environ["DRUPAL_CONTENT_TYPE"]

    spreadsheet = args["spreadsheet"]
    if not spreadsheet.exists():
        logger.error("Spreadsheet not found: %s", spreadsheet)
        sys.exit(1)

    logger.info("Scanning '%s' for rows with Business Criticality Level = 'Core Infrastructure'...", spreadsheet)
    titles = list(iter_core_infrastructure_titles(spreadsheet))
    logger.info("Found %d unique node(s) to patch", len(titles))

    if not titles:
        logger.info("Nothing to do.")
        return

    if args["dry_run"]:
        logger.info("DRY RUN — no nodes will be patched")
        for title in titles:
            logger.info(
                "  Would set field_business_criticality_level='%s' on '%s'",
                PATCH_VALUE,
                title,
            )
        return

    session = requests.Session()
    session.auth = (username, password)
    session.headers.update({
        "Accept": "application/vnd.api+json",
        "Content-Type": "application/vnd.api+json",
    })
    session.verify = False

    logger.info("Fetching all existing nodes from Drupal...")
    node_index = fetch_all_nodes(session, base_url, content_type)
    logger.info("Indexed %d nodes", len(node_index))

    total = len(titles)
    patched = 0
    not_found = 0
    failed = 0
    start_time = time.monotonic()

    for i, title in enumerate(titles, 1):
        uuid = node_index.get(title)
        if not uuid:
            logger.warning("[%d/%d] Node not found in Drupal: '%s'", i, total, title)
            not_found += 1
            continue

        logger.info(
            "[%d/%d] Patching '%s' → field_business_criticality_level='%s'",
            i, total, title, PATCH_VALUE,
        )
        if patch_node(session, base_url, content_type, uuid, title):
            patched += 1
        else:
            failed += 1

    elapsed = time.monotonic() - start_time
    logger.info("=" * 60)
    logger.info("PATCH CORE INFRASTRUCTURE SUMMARY")
    logger.info("  Total entries    : %d", total)
    logger.info("  Patched          : %d", patched)
    logger.info("  Not found        : %d", not_found)
    logger.info("  Failed           : %d", failed)
    logger.info("  Elapsed          : %.1fs", elapsed)


if __name__ == "__main__":
    main()
